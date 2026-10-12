import os
import time
import subprocess
import json
import re
import socket
import urllib.request
import urllib.error
from datetime import datetime

import psutil
import win32gui
import win32process

from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# KONFIGURASI
# ============================================================

PROCESS_NAME = "OnAir4.7.3.8.exe"

BASE_DIR = r"D:\Monitoring_OnAir"

EXCEL_FILE = os.path.join(
    BASE_DIR,
    "Monitoring_OnAir.xlsx"
)

LOG_FILE = os.path.join(
    BASE_DIR,
    "monitoring.log"
)

CHECK_INTERVAL = 5

AGENT_CONFIG_FILE = os.path.join(
    BASE_DIR,
    "agent_config.json"
)

HEARTBEAT_TIMEOUT = 10

DOCKER_COMMAND = "docker"
AGENT_API_URL = os.getenv("AGENT_API_URL", "")
CLIENT_ID = os.getenv("CLIENT_ID", "")
AGENT_TOKEN = os.getenv("AGENT_TOKEN", "")
AGENT_VERSION = os.getenv("AGENT_VERSION", "1.0.0")

# ------------------------------------------------------------
# Docker History
# ------------------------------------------------------------
# Docker tidak akan ditulis ke Excel setiap interval.
#
# Data Docker hanya ditulis jika:
# 1. Status Docker berubah
# 2. Container bertambah / berkurang
# 3. CPU berubah >= threshold
# 4. Memory berubah >= threshold
#
# Contoh:
# CPU sebelumnya 10%, sekarang 25%
# -> akan dicatat
#
# CPU sebelumnya 10%, sekarang 12%
# -> tidak dicatat
# ------------------------------------------------------------

DOCKER_CPU_CHANGE_THRESHOLD = 10.0
DOCKER_MEMORY_CHANGE_THRESHOLD = 10.0


# ============================================================
# FOLDER
# ============================================================

os.makedirs(
    BASE_DIR,
    exist_ok=True
)


# ============================================================
# WARNA EXCEL
# ============================================================

COLOR_HEADER = "1F4E78"
COLOR_HEADER_DARK = "17365D"

COLOR_GREEN = "E2F0D9"
COLOR_RED = "F4CCCC"
COLOR_YELLOW = "FFF2CC"
COLOR_BLUE = "DDEBF7"
COLOR_GRAY = "F2F2F2"

COLOR_BORDER = "D9E2F3"


# ============================================================
# KONSTANTA
# ============================================================

NOT_READABLE = "Tidak Terbaca"

TABLE_COUNTER = 0


# ============================================================
# LOG
# ============================================================

def write_log(message):

    try:

        timestamp = datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )

        with open(
            LOG_FILE,
            "a",
            encoding="utf-8"
        ) as file:

            file.write(
                f"[{timestamp}] {message}\n"
            )

    except Exception:
        pass


# ============================================================
# FORMAT WAKTU
# ============================================================

def format_time(value):

    if value is None:
        return ""

    return value.strftime(
        "%Y-%m-%d %H:%M:%S"
    )


# ============================================================
# BYTE
# ============================================================

def gb(value):

    return value / (
        1024 ** 3
    )


def mbps(value):

    return value / (
        1024 ** 2
    )


# ============================================================
# SAFE FLOAT
# ============================================================

def safe_float(value):

    try:

        if value is None:
            return None

        text = str(value).strip()

        text = (
            text
            .replace("%", "")
            .replace("GB", "")
            .replace("MB", "")
            .replace("KB", "")
            .strip()
        )

        return float(text)

    except Exception:

        return None


# ============================================================
# DISK CAPACITY
# ============================================================

def get_disk_capacity(drive):

    path = drive + ":\\"

    try:

        usage = psutil.disk_usage(
            path
        )

        return {

            "percent":
                round(
                    usage.percent,
                    2
                ),

            "used":
                round(
                    gb(usage.used),
                    2
                ),

            "total":
                round(
                    gb(usage.total),
                    2
                )

        }

    except Exception:

        return {

            "percent": NOT_READABLE,
            "used": NOT_READABLE,
            "total": NOT_READABLE

        }


# ============================================================
# WINDOWS DISK PERFORMANCE
# ============================================================

def get_disk_performance():

    result = {

        "C_active": NOT_READABLE,
        "C_read": NOT_READABLE,
        "C_write": NOT_READABLE,

        "D_active": NOT_READABLE,
        "D_read": NOT_READABLE,
        "D_write": NOT_READABLE

    }

    powershell_script = r"""
$items = Get-Counter '\PhysicalDisk(*)\% Disk Time',
                    '\PhysicalDisk(*)\Disk Read Bytes/sec',
                    '\PhysicalDisk(*)\Disk Write Bytes/sec'

$instance = $items.CounterSamples |
    Group-Object InstanceName

foreach ($group in $instance) {

    $name = $group.Name

    $active = 0
    $read = 0
    $write = 0

    foreach ($sample in $group.Group) {

        if ($sample.Path -like '*% Disk Time') {
            $active = $sample.CookedValue
        }

        elseif ($sample.Path -like '*Disk Read Bytes/sec') {
            $read = $sample.CookedValue
        }

        elseif ($sample.Path -like '*Disk Write Bytes/sec') {
            $write = $sample.CookedValue
        }
    }

    Write-Output "$name|$active|$read|$write"
}
"""

    try:

        completed = subprocess.run(

            [
                "powershell.exe",
                "-NoProfile",
                "-NonInteractive",
                "-Command",
                powershell_script
            ],

            capture_output=True,
            text=True,
            timeout=5

        )

        output = completed.stdout.strip()

        if not output:
            return result

        for line in output.splitlines():

            parts = line.strip().split("|")

            if len(parts) != 4:
                continue

            instance = parts[0]

            try:

                active = float(
                    parts[1]
                )

                read = mbps(
                    float(parts[2])
                )

                write = mbps(
                    float(parts[3])
                )

            except ValueError:

                continue

            instance_upper = (
                instance.upper()
            )

            if "C:" in instance_upper:

                result["C_active"] = round(
                    active,
                    2
                )

                result["C_read"] = round(
                    read,
                    2
                )

                result["C_write"] = round(
                    write,
                    2
                )

            elif "D:" in instance_upper:

                result["D_active"] = round(
                    active,
                    2
                )

                result["D_read"] = round(
                    read,
                    2
                )

                result["D_write"] = round(
                    write,
                    2
                )

    except Exception as error:

        write_log(
            f"Disk Performance error: {error}"
        )

    return result


# ============================================================
# CPU + RAM ON AIR
# ============================================================

def get_onair_resource(process):

    try:

        cpu = process.cpu_percent(
            interval=0.3
        )

    except Exception:

        cpu = NOT_READABLE


    try:

        memory = process.memory_info()

        ram_used = gb(
            memory.rss
        )

    except Exception:

        ram_used = NOT_READABLE


    return {

        "cpu":
            round(cpu, 2)
            if isinstance(
                cpu,
                (int, float)
            )
            else cpu,

        "ram":
            round(ram_used, 2)
            if isinstance(
                ram_used,
                (int, float)
            )
            else ram_used

    }


# ============================================================
# DOCKER MONITORING
# ============================================================

def get_docker_stats():
    """
    Mengambil data container Docker dengan dua perintah terpisah:
    - docker ps     -> ID, nama, image, status
    - docker stats  -> CPU, memory, network, block I/O, PIDs

    Tidak menggunakan .Image pada docker stats karena field tersebut
    memang tidak tersedia pada template stats di versi Docker tertentu.
    """

    containers = []

    try:
        # --------------------------------------------------------
        # 1. Ambil metadata container dari docker ps
        # --------------------------------------------------------
        ps_completed = subprocess.run(
            [
                DOCKER_COMMAND,
                "ps",
                "--format",
                "{{.ID}}|{{.Names}}|{{.Image}}|{{.Status}}"
            ],
            capture_output=True,
            text=True,
            timeout=10
        )

        if ps_completed.returncode != 0:
            error_message = ps_completed.stderr.strip()
            return {
                "available": False,
                "error": error_message or "Tidak dapat membaca container Docker.",
                "containers": []
            }

        metadata = {}

        for line in ps_completed.stdout.strip().splitlines():
            parts = line.split("|", 3)
            if len(parts) != 4:
                continue

            container_id = parts[0].strip()
            name = parts[1].strip()
            image = parts[2].strip()
            status = parts[3].strip()

            image_repository = image
            image_tag = ""

            if "@" in image:
                image_repository = image.split("@", 1)[0]
                image_tag = image.split("@", 1)[1]
            elif ":" in image.rsplit("/", 1)[-1]:
                image_repository, image_tag = image.rsplit(":", 1)

            aircast_version = None
            version_match = re.search(
                r"(\d+\.\d+\.\d+(?:[-+._][0-9A-Za-z.-]+)?)",
                image_tag
            )
            if version_match:
                aircast_version = version_match.group(1)

            metadata[container_id] = {
                "id": container_id,
                "name": name,
                "image": image,
                "image_repository": image_repository,
                "image_tag": image_tag,
                "aircast_version": aircast_version,
                "status": status,
            }

        # --------------------------------------------------------
        # 2. Ambil resource stats dari docker stats
        # --------------------------------------------------------
        stats_format = (
            "{{.ID}}|"
            "{{.Name}}|"
            "{{.CPUPerc}}|"
            "{{.MemUsage}}|"
            "{{.MemPerc}}|"
            "{{.NetIO}}|"
            "{{.BlockIO}}|"
            "{{.PIDs}}"
        )

        stats_completed = subprocess.run(
            [
                DOCKER_COMMAND,
                "stats",
                "--no-stream",
                "--format",
                stats_format
            ],
            capture_output=True,
            text=True,
            timeout=10
        )

        if stats_completed.returncode != 0:
            error_message = stats_completed.stderr.strip()
            return {
                "available": False,
                "error": error_message or "Docker stats tidak tersedia.",
                "containers": []
            }

        # --------------------------------------------------------
        # 3. Gabungkan metadata + stats berdasarkan container ID
        # --------------------------------------------------------
        for line in stats_completed.stdout.strip().splitlines():
            parts = line.split("|", 7)
            if len(parts) != 8:
                continue

            container_id = parts[0].strip()
            name = parts[1].strip()
            cpu = parts[2].strip()
            mem_usage = parts[3].strip()
            mem_percent = parts[4].strip()
            net_io = parts[5].strip()
            block_io = parts[6].strip()
            pids = parts[7].strip()

            info = metadata.get(container_id, {})

            image = info.get("image", "")
            image_repository = info.get("image_repository", image)
            image_tag = info.get("image_tag", "")
            aircast_version = info.get("aircast_version")
            status = info.get("status", "running")

            memory_used = NOT_READABLE
            memory_limit = NOT_READABLE
            if "/" in mem_usage:
                memory_parts = mem_usage.split("/", 1)
                memory_used = memory_parts[0].strip()
                memory_limit = memory_parts[1].strip()

            network_rx = NOT_READABLE
            network_tx = NOT_READABLE
            if "/" in net_io:
                network_parts = net_io.split("/", 1)
                network_rx = network_parts[0].strip()
                network_tx = network_parts[1].strip()

            block_read = NOT_READABLE
            block_write = NOT_READABLE
            if "/" in block_io:
                block_parts = block_io.split("/", 1)
                block_read = block_parts[0].strip()
                block_write = block_parts[1].strip()

            containers.append({
                "id": container_id,
                "name": info.get("name", name),
                "image": image,
                "image_repository": image_repository,
                "image_tag": image_tag,
                "aircast_version": aircast_version,
                "status": status,
                "cpu": cpu,
                "cpu_value": safe_float(cpu),
                "mem_usage": mem_usage,
                "memory_used": memory_used,
                "memory_limit": memory_limit,
                "mem_percent": mem_percent,
                "mem_percent_value": safe_float(mem_percent),
                "net_io": net_io,
                "network_rx": network_rx,
                "network_tx": network_tx,
                "block_io": block_io,
                "block_read": block_read,
                "block_write": block_write,
                "pids": pids
            })

        return {
            "available": True,
            "error": "",
            "containers": containers
        }

    except FileNotFoundError:
        return {
            "available": False,
            "error": (
                "Perintah Docker tidak ditemukan. "
                "Pastikan Docker Desktop sudah terinstall "
                "dan docker.exe masuk PATH."
            ),
            "containers": []
        }

    except subprocess.TimeoutExpired:
        return {
            "available": False,
            "error": "Docker command timeout.",
            "containers": []
        }

    except Exception as error:
        return {
            "available": False,
            "error": str(error),
            "containers": []
        }


# ============================================================
# AIRCAST IMAGE VERSION
# ============================================================

def get_aircast_images():
    """
    Mengambil seluruh image lokal yang repository-nya mengandung
    kata 'aircast'.

    Contoh:
        registry.svara.id/aircast/aircast-app:2.0.15
        registry.svara.id/aircast/aircast-gosrv:latest
        registry.svara.id/aircast/aircast-db:latest
    """

    images = []

    image_format = (
        "{{.Repository}}|"
        "{{.Tag}}|"
        "{{.ID}}|"
        "{{.CreatedAt}}|"
        "{{.Size}}"
    )

    try:
        completed = subprocess.run(
            [
                DOCKER_COMMAND,
                "images",
                "--format",
                image_format
            ],
            capture_output=True,
            text=True,
            timeout=10
        )

        if completed.returncode != 0:
            return {
                "available": False,
                "error": completed.stderr.strip()
                    or "Tidak dapat membaca image container runtime.",
                "images": []
            }

        for line in completed.stdout.strip().splitlines():

            parts = line.split("|")

            if len(parts) != 5:
                continue

            repository = parts[0].strip()
            tag = parts[1].strip()
            image_id = parts[2].strip()
            created = parts[3].strip()
            size = parts[4].strip()

            # Hanya ambil repository AirCast.
            if "aircast" not in repository.lower():
                continue

            version = None

            version_match = re.search(
                r"(\d+\.\d+\.\d+(?:[-+._][0-9A-Za-z.-]+)?)",
                tag
            )

            if version_match:
                version = version_match.group(1)

            images.append({
                "repository": repository,
                "tag": tag,
                "image": f"{repository}:{tag}",
                "version": version,
                "id": image_id,
                "created": created,
                "size": size
            })

        return {
            "available": True,
            "error": "",
            "images": images
        }

    except FileNotFoundError:
        return {
            "available": False,
            "error": (
                f"Perintah {DOCKER_COMMAND} tidak ditemukan."
            ),
            "images": []
        }

    except subprocess.TimeoutExpired:
        return {
            "available": False,
            "error": "Pembacaan AirCast image timeout.",
            "images": []
        }

    except Exception as error:
        return {
            "available": False,
            "error": str(error),
            "images": []
        }


# ============================================================
# WRITE STATUS JSON
# ============================================================

def write_status_json(
    current_time,
    status,
    info,
    docker_info,
    aircast_images
):
    """
    Menulis snapshot terbaru ke JSON agar mudah dibaca website.
    """

    try:

        payload = {
            "timestamp": format_time(current_time),
            "status": status,

            "system": {
                "cpu_percent": info["cpu"],
                "ram_percent": info["ram_percent"],
                "ram_used_gb": info["ram_used"],
                "ram_total_gb": info["ram_total"]
            },

            "onair": {
                "process": PROCESS_NAME,
                "cpu_percent": info["onair_cpu"],
                "ram_gb": info["onair_ram"]
            },

            "disk": {
                "C": {
                    "usage_percent": info["disk_c_percent"],
                    "used_gb": info["disk_c_used"],
                    "total_gb": info["disk_c_total"],
                    "active_percent": info["disk_c_active"],
                    "read_mbps": info["disk_c_read"],
                    "write_mbps": info["disk_c_write"]
                },
                "D": {
                    "usage_percent": info["disk_d_percent"],
                    "used_gb": info["disk_d_used"],
                    "total_gb": info["disk_d_total"],
                    "active_percent": info["disk_d_active"],
                    "read_mbps": info["disk_d_read"],
                    "write_mbps": info["disk_d_write"]
                }
            },

            "container_runtime": {
                "command": DOCKER_COMMAND,
                "available": docker_info["available"],
                "error": docker_info["error"],
                "container_count": len(
                    docker_info["containers"]
                ) if docker_info["available"] else 0,

                "containers": docker_info["containers"]
            },

            "aircast_images": aircast_images,

            "agent": {
                "check_interval_seconds": CHECK_INTERVAL
            }
        }

        temp_file = STATUS_JSON_FILE + ".tmp"

        with open(
            temp_file,
            "w",
            encoding="utf-8"
        ) as file:

            json.dump(
                payload,
                file,
                indent=2,
                ensure_ascii=False
            )

        # Atomic-ish replacement supaya website tidak membaca
        # file JSON yang sedang ditulis.
        os.replace(
            temp_file,
            STATUS_JSON_FILE
        )

    except Exception as error:

        write_log(
            f"Status JSON error: {error}"
        )


# ============================================================
# SEND HEARTBEAT TO AIRCAST MONITOR BACKEND
# ============================================================

def send_heartbeat(current_time, process, info, docker_info, aircast_images, status):
    if not AGENT_API_URL or not CLIENT_ID or not AGENT_TOKEN:
        return False

    try:
        containers = docker_info.get("containers", []) if docker_info.get("available") else []
        payload = {
            "client_id": CLIENT_ID,
            "agent_version": AGENT_VERSION,
            "hostname": socket.gethostname(),
            "timestamp": current_time.isoformat(),
            "system": {
                "cpu_usage": info.get("cpu"),
                "ram_usage": info.get("ram_percent"),
                "disk_usage": info.get("disk_c_percent"),
            },
            "onair": {
                "status": "running" if process is not None and status not in ("NOT_RUNNING", "HANG") else "not_running",
                "cpu_usage": info.get("onair_cpu"),
                "ram_usage": info.get("onair_ram"),
            },
            "runtime": {
                "type": DOCKER_COMMAND,
                "available": docker_info.get("available", False),
                "containers": containers,
            },
            "containers": containers,
            "aircast_images": aircast_images,
        }
        request = urllib.request.Request(
            AGENT_API_URL,
            data=json.dumps(payload).encode("utf-8"),
            headers={
                "Content-Type": "application/json",
                "Authorization": f"Bearer {AGENT_TOKEN}",
                "X-Client-ID": CLIENT_ID,
            },
            method="POST",
        )
        with urllib.request.urlopen(request, timeout=10) as response:
            return 200 <= response.status < 300
    except Exception as error:
        write_log(f"Heartbeat gagal: {error}")
        return False


# ============================================================
# GET SYSTEM INFO
# ============================================================

def get_system_info(process=None):

    # ========================================================
    # CPU
    # ========================================================

    try:

        cpu = psutil.cpu_percent(
            interval=0.3
        )

    except Exception:

        cpu = NOT_READABLE


    # ========================================================
    # RAM
    # ========================================================

    try:

        ram = psutil.virtual_memory()

        ram_percent = round(
            ram.percent,
            2
        )

        ram_used = round(
            gb(ram.used),
            2
        )

        ram_total = round(
            gb(ram.total),
            2
        )

    except Exception:

        ram_percent = NOT_READABLE
        ram_used = NOT_READABLE
        ram_total = NOT_READABLE


    # ========================================================
    # DISK
    # ========================================================

    disk_c = get_disk_capacity(
        "C"
    )

    disk_d = get_disk_capacity(
        "D"
    )


    # ========================================================
    # DISK PERFORMANCE
    # ========================================================

    disk_perf = (
        get_disk_performance()
    )


    # ========================================================
    # ON AIR
    # ========================================================

    if process is not None:

        onair = (
            get_onair_resource(
                process
            )
        )

    else:

        onair = {

            "cpu": NOT_READABLE,
            "ram": NOT_READABLE

        }


    return {

        "cpu":
            round(cpu, 2)
            if isinstance(
                cpu,
                (int, float)
            )
            else cpu,

        "ram_percent":
            ram_percent,

        "ram_used":
            ram_used,

        "ram_total":
            ram_total,

        "onair_cpu":
            onair["cpu"],

        "onair_ram":
            onair["ram"],

        "disk_c_percent":
            disk_c["percent"],

        "disk_c_used":
            disk_c["used"],

        "disk_c_total":
            disk_c["total"],

        "disk_d_percent":
            disk_d["percent"],

        "disk_d_used":
            disk_d["used"],

        "disk_d_total":
            disk_d["total"],

        "disk_c_active":
            disk_perf["C_active"],

        "disk_c_read":
            disk_perf["C_read"],

        "disk_c_write":
            disk_perf["C_write"],

        "disk_d_active":
            disk_perf["D_active"],

        "disk_d_read":
            disk_perf["D_read"],

        "disk_d_write":
            disk_perf["D_write"]

    }


# ============================================================
# FIND ON AIR
# ============================================================

def find_onair():

    for process in psutil.process_iter(
        ["pid", "name"]
    ):

        try:

            name = process.info[
                "name"
            ]

            if (

                name

                and

                name.lower()
                ==
                PROCESS_NAME.lower()

            ):

                return process


        except (

            psutil.NoSuchProcess,
            psutil.AccessDenied

        ):

            continue


    return None


# ============================================================
# CHECK WINDOW
# ============================================================

def check_onair_window(pid):

    windows = []


    def enum_callback(
        hwnd,
        extra
    ):

        try:

            _, window_pid = (

                win32process
                .GetWindowThreadProcessId(
                    hwnd
                )

            )

            if (

                window_pid == pid

                and

                win32gui.IsWindowVisible(
                    hwnd
                )

            ):

                windows.append(
                    hwnd
                )


        except Exception:

            pass


        return True


    try:

        win32gui.EnumWindows(
            enum_callback,
            None
        )

    except Exception:

        return "UNKNOWN"


    if not windows:

        return "RUNNING_NO_WINDOW"


    for hwnd in windows:

        try:

            if win32gui.IsHungAppWindow(
                hwnd
            ):

                return "HANG"

        except Exception:

            pass


    return "NORMAL"


# ============================================================
# EXCEL BORDER
# ============================================================

def get_border():

    thin = Side(
        style="thin",
        color=COLOR_BORDER
    )

    return Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )


# ============================================================
# STYLE HEADER
# ============================================================

def style_header(ws):

    border = get_border()

    for cell in ws[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF",
            size=11
        )

        cell.fill = PatternFill(
            "solid",
            fgColor=COLOR_HEADER
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = border


    ws.row_dimensions[
        1
    ].height = 55


    ws.freeze_panes = "A2"


# ============================================================
# STYLE DATA
# ============================================================

def style_data(ws):

    border = get_border()

    for row in ws.iter_rows(
        min_row=2
    ):

        max_text_length = 0

        for cell in row:

            cell.border = border

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

            if cell.value is not None:

                text_length = len(
                    str(cell.value)
                )

                max_text_length = max(
                    max_text_length,
                    text_length
                )

        # ----------------------------------------------------
        # Tinggi baris menyesuaikan isi
        # ----------------------------------------------------

        if max_text_length > 100:

            height = 65

        elif max_text_length > 70:

            height = 55

        elif max_text_length > 45:

            height = 45

        elif max_text_length > 25:

            height = 38

        else:

            height = 30


        ws.row_dimensions[
            row[0].row
        ].height = height


# ============================================================
# AUTO WIDTH
# ============================================================

def auto_fit_columns(
    ws,
    min_width=12,
    max_width=45
):

    for column_cells in ws.columns:

        column_letter = (
            column_cells[0].column_letter
        )

        max_length = 0

        for cell in column_cells:

            try:

                if cell.value is not None:

                    value = str(
                        cell.value
                    )

                    longest_line = max(
                        (
                            len(line)
                            for line
                            in value.split("\n")
                        ),
                        default=0
                    )

                    max_length = max(
                        max_length,
                        longest_line
                    )

            except Exception:

                pass


        width = (
            max_length + 3
        )

        width = max(
            width,
            min_width
        )

        width = min(
            width,
            max_width
        )

        ws.column_dimensions[
            column_letter
        ].width = width


# ============================================================
# TABLE
# ============================================================

def add_excel_table(
    ws,
    table_name
):

    global TABLE_COUNTER

    if ws.max_row < 2:
        return


    for name in list(
        ws.tables.keys()
    ):

        del ws.tables[
            name
        ]


    last_column = (
        ws.cell(
            1,
            ws.max_column
        ).column_letter
    )


    table_ref = (
        f"A1:{last_column}"
        f"{ws.max_row}"
    )


    TABLE_COUNTER += 1


    safe_name = (
        f"{table_name}{TABLE_COUNTER}"
    )


    table = Table(
        displayName=safe_name,
        ref=table_ref
    )


    style = TableStyleInfo(

        name="TableStyleMedium2",

        showFirstColumn=False,

        showLastColumn=False,

        showRowStripes=True,

        showColumnStripes=False

    )


    table.tableStyleInfo = style

    ws.add_table(
        table
    )


# ============================================================
# EVENT COLUMN WIDTH
# ============================================================

def event_column_widths(ws):

    widths = {

        "A": 22,
        "B": 22,
        "C": 28,
        "D": 18,
        "E": 12,

        "F": 19,
        "G": 19,

        "H": 19,
        "I": 20,
        "J": 20,

        "K": 20,

        "L": 21,
        "M": 21,
        "N": 21,

        "O": 21,
        "P": 21,
        "Q": 21,

        "R": 25,
        "S": 23,
        "T": 23,

        "U": 25,
        "V": 23,
        "W": 23

    }


    for column, width in widths.items():

        ws.column_dimensions[
            column
        ].width = width


# ============================================================
# RESOURCE COLUMN WIDTH
# ============================================================

def resource_column_widths(ws):

    widths = {

        "A": 22,
        "B": 24,

        "C": 20,
        "D": 20,

        "E": 20,
        "F": 20,
        "G": 20,

        "H": 20,

        "I": 22,
        "J": 22,
        "K": 22,

        "L": 22,
        "M": 22,
        "N": 22,

        "O": 22,
        "P": 22,
        "Q": 22,

        "R": 22,
        "S": 22,
        "T": 22

    }


    for column, width in widths.items():

        ws.column_dimensions[
            column
        ].width = width


# ============================================================
# DOCKER COLUMN WIDTH
# ============================================================

def docker_column_widths(ws):

    widths = {

        "A": 22,
        "B": 22,
        "C": 28,

        "D": 20,

        "E": 22,
        "F": 22,
        "G": 22,

        "H": 22,

        "I": 20,
        "J": 20,

        "K": 22,
        "L": 22,

        "M": 22,
        "N": 22,

        "O": 12

    }


    for column, width in widths.items():

        ws.column_dimensions[
            column
        ].width = width


# ============================================================
# FORMAT EVENT SHEET
# ============================================================

def format_event_sheet(ws):

    style_header(
        ws
    )

    style_data(
        ws
    )

    event_column_widths(
        ws
    )


    for row in range(
        2,
        ws.max_row + 1
    ):

        ws.cell(
            row,
            1
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )

        ws.cell(
            row,
            2
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )


    for row in range(
        2,
        ws.max_row + 1
    ):

        for col in range(
            4,
            24
        ):

            value = ws.cell(
                row,
                col
            ).value

            if isinstance(
                value,
                (int, float)
            ):

                ws.cell(
                    row,
                    col
                ).number_format = (
                    "0.00"
                )


    for row in range(
        2,
        ws.max_row + 1
    ):

        event = str(
            ws.cell(
                row,
                3
            ).value or ""
        ).upper()


        if event == "START":

            fill = PatternFill(
                "solid",
                fgColor=COLOR_GREEN
            )

        elif event == "HANG":

            fill = PatternFill(
                "solid",
                fgColor=COLOR_RED
            )

        elif "RECOVERED" in event:

            fill = PatternFill(
                "solid",
                fgColor=COLOR_BLUE
            )

        elif "EXITED" in event:

            fill = PatternFill(
                "solid",
                fgColor=COLOR_YELLOW
            )

        else:

            fill = PatternFill(
                "solid",
                fgColor="FFFFFF"
            )


        for col in range(
            1,
            ws.max_column + 1
        ):

            ws.cell(
                row,
                col
            ).fill = fill


# ============================================================
# FORMAT RESOURCE SHEET
# ============================================================

def format_resource_sheet(ws):

    style_header(
        ws
    )

    style_data(
        ws
    )

    resource_column_widths(
        ws
    )


    for row in range(
        2,
        ws.max_row + 1
    ):

        ws.cell(
            row,
            1
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )


        for col in range(
            3,
            21
        ):

            value = ws.cell(
                row,
                col
            ).value

            if isinstance(
                value,
                (int, float)
            ):

                ws.cell(
                    row,
                    col
                ).number_format = (
                    "0.00"
                )


# ============================================================
# FORMAT DOCKER SHEET
# ============================================================

def format_docker_sheet(ws):

    style_header(
        ws
    )

    style_data(
        ws
    )

    docker_column_widths(
        ws
    )


    for row in range(
        2,
        ws.max_row + 1
    ):

        ws.cell(
            row,
            1
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )


        for col in range(
            4,
            16
        ):

            ws.cell(
                row,
                col
            ).alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )


# ============================================================
# CREATE EXCEL
# ============================================================

def create_excel():

    if os.path.exists(
        EXCEL_FILE
    ):

        return


    wb = Workbook()


    # ========================================================
    # DASHBOARD
    # ========================================================

    dashboard = wb.active

    dashboard.title = (
        "Dashboard"
    )


    dashboard["A1"] = (
        "MONITORING PLAYER ON AIR + DOCKER"
    )


    dashboard["A1"].font = Font(
        bold=True,
        size=20,
        color="FFFFFF"
    )


    dashboard["A1"].fill = PatternFill(
        "solid",
        fgColor=COLOR_HEADER_DARK
    )


    dashboard.merge_cells(
        "A1:F1"
    )


    dashboard["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


    dashboard.row_dimensions[
        1
    ].height = 45


    dashboard["A3"] = "Process"
    dashboard["B3"] = PROCESS_NAME

    dashboard["A4"] = "Interval"
    dashboard["B4"] = (
        f"{CHECK_INTERVAL} detik"
    )

    dashboard["A5"] = "Folder"
    dashboard["B5"] = BASE_DIR

    dashboard["A6"] = "Excel"
    dashboard["B6"] = EXCEL_FILE

    dashboard["A7"] = "Log"
    dashboard["B7"] = LOG_FILE


    for row in range(
        3,
        8
    ):

        dashboard.cell(
            row,
            1
        ).font = Font(
            bold=True
        )

        dashboard.cell(
            row,
            1
        ).fill = PatternFill(
            "solid",
            fgColor=COLOR_BLUE
        )

        dashboard.cell(
            row,
            1
        ).border = get_border()

        dashboard.cell(
            row,
            2
        ).border = get_border()

        dashboard.cell(
            row,
            2
        ).alignment = Alignment(
            vertical="center",
            wrap_text=True
        )


    dashboard.column_dimensions[
        "A"
    ].width = 28

    dashboard.column_dimensions[
        "B"
    ].width = 80

    dashboard.column_dimensions[
        "C"
    ].width = 5

    dashboard.column_dimensions[
        "D"
    ].width = 25

    dashboard.column_dimensions[
        "E"
    ].width = 28

    dashboard.column_dimensions[
        "F"
    ].width = 5


    # ========================================================
    # EVENT LOG
    # ========================================================

    event_ws = wb.create_sheet(
        "Event Log"
    )


    event_headers = [

        "Waktu Mulai",
        "Waktu Selesai",
        "Event",
        "Durasi (detik)",
        "PID",

        "CPU System (%)",
        "CPU OnAir (%)",

        "RAM System (%)",
        "RAM Used (GB)",
        "RAM Total (GB)",

        "RAM OnAir (GB)",

        "Disk C Usage (%)",
        "Disk C Used (GB)",
        "Disk C Total (GB)",

        "Disk D Usage (%)",
        "Disk D Used (GB)",
        "Disk D Total (GB)",

        "Disk C Active Time (%)",
        "Disk C Read (MB/s)",
        "Disk C Write (MB/s)",

        "Disk D Active Time (%)",
        "Disk D Read (MB/s)",
        "Disk D Write (MB/s)"

    ]


    event_ws.append(
        event_headers
    )


    # ========================================================
    # RESOURCE HISTORY
    # ========================================================

    resource_ws = wb.create_sheet(
        "Resource History"
    )


    resource_headers = [

        "Waktu",

        "Status",

        "CPU System (%)",

        "CPU OnAir (%)",

        "RAM System (%)",

        "RAM Used (GB)",

        "RAM Total (GB)",

        "RAM OnAir (GB)",

        "Disk C Usage (%)",

        "Disk C Used (GB)",

        "Disk C Total (GB)",

        "Disk C Active (%)",

        "Disk C Read (MB/s)",

        "Disk C Write (MB/s)",

        "Disk D Usage (%)",

        "Disk D Used (GB)",

        "Disk D Total (GB)",

        "Disk D Active (%)",

        "Disk D Read (MB/s)",

        "Disk D Write (MB/s)"

    ]


    resource_ws.append(
        resource_headers
    )


    # ========================================================
    # DOCKER HISTORY
    # ========================================================

    docker_ws = wb.create_sheet(
        "Docker History"
    )


    docker_headers = [

        "Waktu",

        "Status OnAir",

        "Container",

        "Container CPU (%)",

        "Memory Usage",

        "Memory Used",

        "Memory Limit",

        "Memory (%)",

        "Network I/O",

        "Network RX",

        "Network TX",

        "Block I/O",

        "Block Read",

        "Block Write",

        "PIDs"

    ]


    docker_ws.append(
        docker_headers
    )


    # ========================================================
    # INFO
    # ========================================================

    info = wb.create_sheet(
        "Info"
    )


    info["A1"] = (
        "INFORMASI MONITORING"
    )


    info["A1"].font = Font(
        bold=True,
        size=18,
        color="FFFFFF"
    )


    info["A1"].fill = PatternFill(
        "solid",
        fgColor=COLOR_HEADER_DARK
    )


    info.merge_cells(
        "A1:B1"
    )


    info_data = [

        (
            "Process",
            PROCESS_NAME
        ),

        (
            "Interval",
            f"{CHECK_INTERVAL} detik"
        ),

        (
            "Monitoring Windows",
            "CPU, RAM, Disk Capacity, "
            "Disk Active Time, Read, Write"
        ),

        (
            "Monitoring OnAir",
            "CPU, RAM, Window Status"
        ),

        (
            "Monitoring Docker",
            "Container, CPU, Memory, "
            "Memory %, Network I/O, "
            "Block I/O, PIDs"
        ),

        (
            "Docker History",
            "Hanya disimpan jika terjadi "
            "perubahan penting pada Docker"
        ),

        (
            "Docker CPU Threshold",
            f"{DOCKER_CPU_CHANGE_THRESHOLD}%"
        ),

        (
            "Docker Memory Threshold",
            f"{DOCKER_MEMORY_CHANGE_THRESHOLD}%"
        ),

        (
            "Event Log",
            "START, HANG, RECOVERED, "
            "PROCESS EXITED"
        ),

        (
            "Resource History",
            "Data resource komputer "
            "setiap interval monitoring"
        ),

        (
            "Log File",
            LOG_FILE
        )

    ]


    row = 3

    for title, value in info_data:

        info.cell(
            row,
            1
        ).value = title

        info.cell(
            row,
            2
        ).value = value

        info.cell(
            row,
            1
        ).font = Font(
            bold=True
        )

        info.cell(
            row,
            1
        ).fill = PatternFill(
            "solid",
            fgColor=COLOR_BLUE
        )

        info.cell(
            row,
            1
        ).border = get_border()

        info.cell(
            row,
            2
        ).border = get_border()

        info.cell(
            row,
            2
        ).alignment = Alignment(
            vertical="center",
            wrap_text=True
        )

        row += 1


    info.column_dimensions[
        "A"
    ].width = 32

    info.column_dimensions[
        "B"
    ].width = 100


    # ========================================================
    # STYLE
    # ========================================================

    format_event_sheet(
        event_ws
    )

    format_resource_sheet(
        resource_ws
    )

    format_docker_sheet(
        docker_ws
    )


    info.freeze_panes = "A3"

    dashboard.freeze_panes = "A3"


    # ========================================================
    # SAVE
    # ========================================================

    wb.save(
        EXCEL_FILE
    )

    wb.close()


# ============================================================
# ADD EVENT
# ============================================================

def add_event(

    start_time,
    end_time,
    event,
    duration,
    pid,
    info

):

    try:

        wb = load_workbook(
            EXCEL_FILE
        )

        ws = wb[
            "Event Log"
        ]


        ws.append([

            format_time(
                start_time
            ),

            format_time(
                end_time
            ),

            event,

            duration,

            pid,

            info["cpu"],

            info["onair_cpu"],

            info["ram_percent"],

            info["ram_used"],

            info["ram_total"],

            info["onair_ram"],

            info["disk_c_percent"],

            info["disk_c_used"],

            info["disk_c_total"],

            info["disk_d_percent"],

            info["disk_d_used"],

            info["disk_d_total"],

            info["disk_c_active"],

            info["disk_c_read"],

            info["disk_c_write"],

            info["disk_d_active"],

            info["disk_d_read"],

            info["disk_d_write"]

        ])


        format_event_sheet(
            ws
        )


        add_excel_table(
            ws,
            "EventLog"
        )


        wb.save(
            EXCEL_FILE
        )

        wb.close()


    except PermissionError:

        write_log(
            "Excel sedang terbuka. "
            "Event belum dapat disimpan."
        )


    except Exception as error:

        write_log(
            f"Gagal menyimpan event: {error}"
        )


# ============================================================
# UPDATE LAST EVENT
# ============================================================

def update_last_event(

    end_time,
    duration,
    event_name

):

    try:

        wb = load_workbook(
            EXCEL_FILE
        )

        ws = wb[
            "Event Log"
        ]


        last_row = (
            ws.max_row
        )


        if last_row < 2:

            wb.close()

            return


        ws.cell(
            last_row,
            2
        ).value = format_time(
            end_time
        )


        ws.cell(
            last_row,
            3
        ).value = event_name


        ws.cell(
            last_row,
            4
        ).value = duration


        format_event_sheet(
            ws
        )


        add_excel_table(
            ws,
            "EventLog"
        )


        wb.save(
            EXCEL_FILE
        )

        wb.close()


    except PermissionError:

        write_log(
            "Excel sedang terbuka. "
            "Recovery belum dapat diperbarui."
        )


    except Exception as error:

        write_log(
            f"Gagal update event: {error}"
        )


# ============================================================
# ADD RESOURCE HISTORY
# ============================================================

def add_resource_history(

    current_time,
    status,
    info

):

    try:

        wb = load_workbook(
            EXCEL_FILE
        )

        ws = wb[
            "Resource History"
        ]


        ws.append([

            format_time(
                current_time
            ),

            status,

            info["cpu"],

            info["onair_cpu"],

            info["ram_percent"],

            info["ram_used"],

            info["ram_total"],

            info["onair_ram"],

            info["disk_c_percent"],

            info["disk_c_used"],

            info["disk_c_total"],

            info["disk_c_active"],

            info["disk_c_read"],

            info["disk_c_write"],

            info["disk_d_percent"],

            info["disk_d_used"],

            info["disk_d_total"],

            info["disk_d_active"],

            info["disk_d_read"],

            info["disk_d_write"]

        ])


        format_resource_sheet(
            ws
        )


        add_excel_table(
            ws,
            "ResourceHistory"
        )


        wb.save(
            EXCEL_FILE
        )

        wb.close()


    except PermissionError:

        write_log(
            "Excel sedang terbuka. "
            "Resource History belum tersimpan."
        )


    except Exception as error:

        write_log(
            f"Gagal menyimpan Resource History: {error}"
        )


# ============================================================
# DOCKER SNAPSHOT
# ============================================================

def create_docker_snapshot(docker_info):

    snapshot = {

        "available":
            docker_info["available"],

        "containers": {}

    }


    if not docker_info["available"]:

        return snapshot


    for container in docker_info["containers"]:

        name = container["name"]

        snapshot["containers"][
            name
        ] = {

            "id":
                container["id"],

            "cpu":
                container["cpu_value"],

            "memory":
                container["mem_percent_value"],

            "memory_used":
                container["memory_used"],

            "memory_limit":
                container["memory_limit"]

        }


    return snapshot


# ============================================================
# CEK PERUBAHAN DOCKER
# ============================================================

def docker_has_significant_change(

    previous_snapshot,
    current_snapshot

):

    # --------------------------------------------------------
    # Belum pernah ada snapshot
    # --------------------------------------------------------

    if previous_snapshot is None:

        return True


    # --------------------------------------------------------
    # Docker ONLINE / OFFLINE berubah
    # --------------------------------------------------------

    if (

        previous_snapshot["available"]
        !=
        current_snapshot["available"]

    ):

        return True


    # --------------------------------------------------------
    # Kalau Docker OFFLINE
    # --------------------------------------------------------

    if not current_snapshot["available"]:

        return False


    previous_containers = (
        previous_snapshot["containers"]
    )

    current_containers = (
        current_snapshot["containers"]
    )


    # --------------------------------------------------------
    # Container bertambah / berkurang
    # --------------------------------------------------------

    if (

        set(previous_containers.keys())
        !=
        set(current_containers.keys())

    ):

        return True


    # --------------------------------------------------------
    # Periksa setiap container
    # --------------------------------------------------------

    for name in current_containers:

        current = (
            current_containers[name]
        )

        previous = (
            previous_containers[name]
        )


        # ----------------------------------------------------
        # CPU
        # ----------------------------------------------------

        current_cpu = (
            current["cpu"]
        )

        previous_cpu = (
            previous["cpu"]
        )

        if (

            current_cpu is not None

            and

            previous_cpu is not None

        ):

            if abs(
                current_cpu
                -
                previous_cpu
            ) >= DOCKER_CPU_CHANGE_THRESHOLD:

                return True


        # ----------------------------------------------------
        # Memory %
        # ----------------------------------------------------

        current_memory = (
            current["memory"]
        )

        previous_memory = (
            previous["memory"]
        )

        if (

            current_memory is not None

            and

            previous_memory is not None

        ):

            if abs(
                current_memory
                -
                previous_memory
            ) >= DOCKER_MEMORY_CHANGE_THRESHOLD:

                return True


        # ----------------------------------------------------
        # Container ID berubah
        # ----------------------------------------------------

        if (

            current["id"]
            !=
            previous["id"]

        ):

            return True


    return False


# ============================================================
# ADD DOCKER HISTORY
# ============================================================

def add_docker_history(

    current_time,
    status,
    docker_info

):

    try:

        wb = load_workbook(
            EXCEL_FILE
        )

        ws = wb[
            "Docker History"
        ]


        # ====================================================
        # DOCKER TIDAK TERSEDIA
        # ====================================================

        if not docker_info["available"]:

            ws.append([

                format_time(
                    current_time
                ),

                status,

                "DOCKER OFFLINE",

                NOT_READABLE,

                docker_info["error"],

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE

            ])


        # ====================================================
        # TIDAK ADA CONTAINER
        # ====================================================

        elif not docker_info["containers"]:

            ws.append([

                format_time(
                    current_time
                ),

                status,

                "NO CONTAINER",

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE,

                NOT_READABLE

            ])


        # ====================================================
        # CONTAINER
        # ====================================================

        else:

            for container in (
                docker_info["containers"]
            ):

                ws.append([

                    format_time(
                        current_time
                    ),

                    status,

                    container["name"],

                    container["cpu"],

                    container["mem_usage"],

                    container["memory_used"],

                    container["memory_limit"],

                    container["mem_percent"],

                    container["net_io"],

                    container["network_rx"],

                    container["network_tx"],

                    container["block_io"],

                    container["block_read"],

                    container["block_write"],

                    container["pids"]

                ])


        format_docker_sheet(
            ws
        )


        add_excel_table(
            ws,
            "DockerHistory"
        )


        wb.save(
            EXCEL_FILE
        )

        wb.close()


    except PermissionError:

        write_log(
            "Excel sedang terbuka. "
            "Docker History belum tersimpan."
        )


    except Exception as error:

        write_log(
            f"Gagal menyimpan Docker History: {error}"
        )


# ============================================================
# UPDATE DASHBOARD
# ============================================================

def update_dashboard(

    current_time,
    status,
    info,
    docker_info

):

    try:

        wb = load_workbook(
            EXCEL_FILE
        )

        ws = wb[
            "Dashboard"
        ]


        # ====================================================
        # STATUS
        # ====================================================

        ws["D3"] = "STATUS ON AIR"

        ws["E3"] = status


        ws["D4"] = "CPU SYSTEM"

        ws["E4"] = info[
            "cpu"
        ]


        ws["D5"] = "CPU ONAIR"

        ws["E5"] = info[
            "onair_cpu"
        ]


        ws["D6"] = "RAM SYSTEM"

        ws["E6"] = info[
            "ram_percent"
        ]


        ws["D7"] = "RAM ONAIR"

        ws["E7"] = info[
            "onair_ram"
        ]


        # ====================================================
        # DOCKER
        # ====================================================

        ws["D9"] = "DOCKER STATUS"

        if docker_info["available"]:

            ws["E9"] = "ONLINE"

            ws["D10"] = "CONTAINER"

            ws["E10"] = len(
                docker_info[
                    "containers"
                ]
            )

        else:

            ws["E9"] = "OFFLINE"

            ws["D10"] = "CONTAINER"

            ws["E10"] = 0


        # ====================================================
        # WAKTU
        # ====================================================

        ws["D12"] = "LAST UPDATE"

        ws["E12"] = format_time(
            current_time
        )


        # ====================================================
        # STYLE
        # ====================================================

        for row in range(
            3,
            13
        ):

            ws.cell(
                row,
                4
            ).font = Font(
                bold=True
            )

            ws.cell(
                row,
                4
            ).fill = PatternFill(
                "solid",
                fgColor=COLOR_BLUE
            )

            ws.cell(
                row,
                4
            ).border = get_border()

            ws.cell(
                row,
                5
            ).border = get_border()

            ws.cell(
                row,
                5
            ).alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )


        # ====================================================
        # WARNA STATUS
        # ====================================================

        status_upper = str(
            status
        ).upper()


        if status_upper == "NORMAL":

            ws["E3"].fill = PatternFill(
                "solid",
                fgColor=COLOR_GREEN
            )

        elif status_upper == "HANG":

            ws["E3"].fill = PatternFill(
                "solid",
                fgColor=COLOR_RED
            )

        elif status_upper == "NOT_RUNNING":

            ws["E3"].fill = PatternFill(
                "solid",
                fgColor=COLOR_YELLOW
            )

        else:

            ws["E3"].fill = PatternFill(
                "solid",
                fgColor=COLOR_GRAY
            )


        # ====================================================
        # DOCKER STATUS COLOR
        # ====================================================

        if docker_info["available"]:

            ws["E9"].fill = PatternFill(
                "solid",
                fgColor=COLOR_GREEN
            )

        else:

            ws["E9"].fill = PatternFill(
                "solid",
                fgColor=COLOR_RED
            )


        wb.save(
            EXCEL_FILE
        )

        wb.close()


    except PermissionError:

        write_log(
            "Excel sedang terbuka. "
            "Dashboard tidak dapat diperbarui."
        )


    except Exception as error:

        write_log(
            f"Dashboard update error: {error}"
        )


# ============================================================
# MAIN
# ============================================================

def main():

    create_excel()


    # ========================================================
    # START LOG
    # ========================================================

    write_log(
        "=============================================="
    )

    write_log(
        "MONITORING PLAYER ON AIR + DOCKER DIMULAI"
    )

    write_log(
        f"Process: {PROCESS_NAME}"
    )

    write_log(
        f"Folder: {BASE_DIR}"
    )

    write_log(
        f"Interval: {CHECK_INTERVAL} detik"
    )

    write_log(
        "Docker History menggunakan change detection."
    )

    write_log(
        "=============================================="
    )


    # ========================================================
    # CONSOLE
    # ========================================================

    print(
        "=============================================="
    )

    print(
        "       MONITORING PLAYER ON AIR"
    )

    print(
        "       WINDOWS + DOCKER"
    )

    print(
        "=============================================="
    )

    print()

    print(
        f"Process : {PROCESS_NAME}"
    )

    print(
        f"Folder  : {BASE_DIR}"
    )

    print(
        f"Excel   : {EXCEL_FILE}"
    )

    print(
        f"Log     : {LOG_FILE}"
    )

    print(
        f"Interval: {CHECK_INTERVAL} detik"
    )

    print()

    print(
        "Docker History : CHANGE DETECTION"
    )

    print(
        f"CPU threshold  : "
        f"{DOCKER_CPU_CHANGE_THRESHOLD}%"
    )

    print(
        f"Memory threshold: "
        f"{DOCKER_MEMORY_CHANGE_THRESHOLD}%"
    )

    print()

    print(
        "Monitoring berjalan..."
    )

    print(
        "Tekan CTRL + C untuk berhenti."
    )

    print()


    previous_status = (
        "NOT_RUNNING"
    )

    current_incident = None

    docker_error_logged = False

    # --------------------------------------------------------
    # Snapshot Docker sebelumnya
    # --------------------------------------------------------

    previous_docker_snapshot = None


    # ========================================================
    # LOOP
    # ========================================================

    while True:

        try:

            current_time = (
                datetime.now()
            )


            # =================================================
            # FIND ON AIR
            # =================================================

            process = find_onair()


            # =================================================
            # DOCKER
            # =================================================

            docker_info = (
                get_docker_stats()
            )

            # =================================================
            # AIRCAST IMAGE
            # =================================================

            aircast_images = (
                get_aircast_images()
            )


            if aircast_images["available"]:
                aircast_versions = [
                    item["version"] or item["tag"]
                    for item in aircast_images["images"]
                ]

                write_log(
                    "AirCast image terdeteksi: "
                    + ", ".join(aircast_versions)
                )


            # =================================================
            # DOCKER CHANGE DETECTION
            # =================================================

            current_docker_snapshot = (
                create_docker_snapshot(
                    docker_info
                )
            )


            docker_changed = (
                docker_has_significant_change(

                    previous_docker_snapshot,

                    current_docker_snapshot

                )
            )


            # =================================================
            # LOG DOCKER ERROR
            # =================================================

            if (

                not docker_info["available"]

                and

                not docker_error_logged

            ):

                write_log(
                    "DOCKER ERROR: "
                    f"{docker_info['error']}"
                )

                docker_error_logged = True


            elif docker_info["available"]:

                docker_error_logged = False


            # =================================================
            # PROCESS TIDAK ADA
            # =================================================

            if process is None:

                current_status = (
                    "NOT_RUNNING"
                )


                # ---------------------------------------------
                # HANG → PROCESS EXITED
                # ---------------------------------------------

                if (

                    previous_status
                    ==
                    "HANG"

                    and

                    current_incident
                    is not None

                ):

                    duration = (

                        current_time
                        -
                        current_incident[
                            "start"
                        ]

                    ).total_seconds()


                    update_last_event(

                        current_time,

                        round(
                            duration,
                            2
                        ),

                        "HANG → PROCESS EXITED"

                    )


                    write_log(
                        "OnAir hilang "
                        "saat HANG."
                    )


                    current_incident = (
                        None
                    )


                # ---------------------------------------------
                # NORMAL → EXITED
                # ---------------------------------------------

                elif previous_status in (

                    "NORMAL",

                    "RUNNING_NO_WINDOW"

                ):

                    info = (
                        get_system_info(
                            None
                        )
                    )


                    add_event(

                        current_time,

                        None,

                        "PROCESS EXITED",

                        None,

                        "",

                        info

                    )


                    write_log(
                        "OnAir process hilang."
                    )


                # ---------------------------------------------
                # RESOURCE
                # ---------------------------------------------

                info = (
                    get_system_info(
                        None
                    )
                )


                add_resource_history(

                    current_time,

                    "NOT_RUNNING",

                    info

                )


                # ---------------------------------------------
                # DOCKER HISTORY
                # ---------------------------------------------
                # Hanya tulis jika berubah.
                # ---------------------------------------------

                if docker_changed:

                    add_docker_history(

                        current_time,

                        "NOT_RUNNING",

                        docker_info

                    )

                    write_log(
                        "Perubahan Docker terdeteksi. "
                        "Docker History diperbarui."
                    )


                update_dashboard(

                    current_time,

                    "NOT_RUNNING",

                    info,

                    docker_info

                )

                write_status_json(
                    current_time,
                    "NOT_RUNNING",
                    info,
                    docker_info,
                    aircast_images
                )

                send_heartbeat(current_time, process, info, docker_info, aircast_images, "NOT_RUNNING")


                previous_status = (
                    current_status
                )


                previous_docker_snapshot = (
                    current_docker_snapshot
                )


                docker_count = 0

                if docker_info["available"]:

                    docker_count = len(
                        docker_info[
                            "containers"
                        ]
                    )


                print(

                    f"[{current_time.strftime('%H:%M:%S')}] "
                    "OnAir: NOT_RUNNING | "
                    f"Docker: "
                    f"{docker_count} "
                    "container | "
                    f"AirCast Image: "
                    f"{len(aircast_images['images']) if aircast_images['available'] else 0}"

                )


                time.sleep(
                    CHECK_INTERVAL
                )

                continue


            # =================================================
            # PROCESS DITEMUKAN
            # =================================================

            pid = process.pid


            current_status = (
                check_onair_window(
                    pid
                )
            )


            # =================================================
            # RESOURCE
            # =================================================

            info = (
                get_system_info(
                    process
                )
            )


            # =================================================
            # RESOURCE HISTORY
            # =================================================

            add_resource_history(

                current_time,

                current_status,

                info

            )


            # =================================================
            # DOCKER HISTORY
            # =================================================
            # Hanya ditulis ketika terjadi perubahan.
            # =================================================

            if docker_changed:

                add_docker_history(

                    current_time,

                    current_status,

                    docker_info

                )

                write_log(
                    "Perubahan Docker terdeteksi. "
                    "Docker History diperbarui."
                )


            # =================================================
            # DASHBOARD
            # =================================================

            update_dashboard(

                current_time,

                current_status,

                info,

                docker_info

            )

            write_status_json(
                current_time,
                current_status,
                info,
                docker_info,
                aircast_images
            )

            send_heartbeat(current_time, process, info, docker_info, aircast_images, current_status)


            # =================================================
            # START
            # =================================================

            if (

                previous_status
                ==
                "NOT_RUNNING"

            ):

                add_event(

                    current_time,

                    None,

                    "START",

                    None,

                    pid,

                    info

                )


                write_log(

                    f"OnAir START. PID={pid}"

                )


            # =================================================
            # NORMAL
            # =================================================

            if current_status in (

                "NORMAL",

                "RUNNING_NO_WINDOW"

            ):

                # ---------------------------------------------
                # HANG → RECOVERED
                # ---------------------------------------------

                if (

                    previous_status
                    ==
                    "HANG"

                    and

                    current_incident
                    is not None

                ):

                    duration = (

                        current_time
                        -
                        current_incident[
                            "start"
                        ]

                    ).total_seconds()


                    update_last_event(

                        current_time,

                        round(
                            duration,
                            2
                        ),

                        "HANG → RECOVERED"

                    )


                    write_log(

                        "OnAir RECOVERED. "

                        f"Durasi HANG "
                        f"{duration:.0f} detik."

                    )


                    current_incident = (
                        None
                    )


                previous_status = (
                    current_status
                )


            # =================================================
            # HANG
            # =================================================

            elif current_status == "HANG":

                # ---------------------------------------------
                # HANG BARU
                # ---------------------------------------------

                if previous_status != "HANG":

                    add_event(

                        current_time,

                        None,

                        "HANG",

                        None,

                        pid,

                        info

                    )


                    current_incident = {

                        "start":
                            current_time,

                        "pid":
                            pid

                    }


                    write_log(
                        "=============================================="
                    )

                    write_log(
                        "!!! ON AIR HANG TERDETEKSI !!!"
                    )

                    write_log(
                        f"PID={pid}"
                    )

                    write_log(
                        f"CPU SYSTEM={info['cpu']}%"
                    )

                    write_log(
                        f"CPU ONAIR={info['onair_cpu']}%"
                    )

                    write_log(
                        f"RAM SYSTEM={info['ram_percent']}%"
                    )

                    write_log(
                        f"RAM ONAIR={info['onair_ram']}GB"
                    )

                    write_log(
                        f"DOCKER CONTAINER="
                        f"{len(docker_info['containers'])}"
                    )

                    write_log(
                        "=============================================="
                    )


                else:

                    write_log(
                        "HANG masih berlangsung."
                    )


                previous_status = (
                    "HANG"
                )


            # =================================================
            # DOCKER COUNT
            # =================================================

            docker_count = 0

            if docker_info["available"]:

                docker_count = len(
                    docker_info[
                        "containers"
                    ]
                )


            # =================================================
            # CONSOLE
            # =================================================

            print(

                f"[{current_time.strftime('%H:%M:%S')}] "

                f"OnAir: {current_status} | "

                f"CPU: {info['cpu']}% | "

                f"CPU OnAir: {info['onair_cpu']}% | "

                f"RAM: {info['ram_percent']}% | "

                f"Docker: {docker_count} container | "

                f"AirCast Image: "
                f"{len(aircast_images['images']) if aircast_images['available'] else 0} | "

                f"Docker History: "
                f"{'UPDATE' if docker_changed else 'NO CHANGE'}"

            )


            # =================================================
            # SIMPAN SNAPSHOT DOCKER
            # =================================================

            previous_docker_snapshot = (
                current_docker_snapshot
            )


            # =================================================
            # WAIT
            # =================================================

            time.sleep(
                CHECK_INTERVAL
            )


        # =====================================================
        # CTRL C
        # =====================================================

        except KeyboardInterrupt:

            write_log(
                "Monitoring dihentikan "
                "secara manual."
            )


            print()

            print(
                "Monitoring dihentikan."
            )

            break


        # =====================================================
        # ERROR
        # =====================================================

        except Exception as error:

            write_log(
                f"ERROR MONITORING: {error}"
            )


            print(
                f"ERROR: {error}"
            )


            time.sleep(
                5
            )


# ============================================================
# START PROGRAM
# ============================================================

if __name__ == "__main__":

    main()